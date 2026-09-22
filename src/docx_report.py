from __future__ import annotations

"""将 Markdown 实验报告转换为 Word 文档。"""

from pathlib import Path


def markdown_to_docx(markdown_text: str, output_path: str | Path) -> None:
    """把常见 Markdown 结构写入 docx。

    这里故意只支持报告会用到的少量语法：标题、列表、普通段落和代码块。
    学习 Demo 阶段先保持可读、可控，后续可以替换成更完整的 Markdown 渲染器。
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is required. Please run: pip install -r requirements.txt") from exc

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)

    in_code_block = False
    code_lines: list[str] = []

    def flush_code_block() -> None:
        nonlocal code_lines
        if not code_lines:
            return
        paragraph = document.add_paragraph()
        run = paragraph.add_run("\n".join(code_lines))
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        code_lines = []

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            if in_code_block:
                flush_code_block()
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if not line:
            continue

        if line.startswith("# "):
            title = document.add_paragraph(style="Title")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.add_run(line[2:].strip())
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line[0:3].replace(".", "").isdigit() and ". " in line[:4]:
            document.add_paragraph(line.split(". ", 1)[1].strip(), style="List Number")
        else:
            document.add_paragraph(line)

    flush_code_block()
    document.save(output)
